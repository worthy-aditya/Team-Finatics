"""Verification tests for the Day 32 authorization appendix."""

import json
from pathlib import Path

from docx import Document

from reports.docx_generator import add_authorization_record_docx
from reports.pdf_generator import SentinelReportPDF
from reports.schema import ActiveFinding, AuthorizationRecord, load_authorization_record


AUTHORIZATION_RECORD = AuthorizationRecord(
    target_domain="http://localhost:3000",
    attestation_user="Suraj Yadav (Tester)",
    attestation_timestamp="2026-09-18T12:00:00Z",
    verification_method=".well-known token check",
    verification_status="VERIFIED",
    token_used="test-token",
)


def test_authorization_record_loads_from_jsonl(tmp_path: Path) -> None:
    log_path = tmp_path / "consent_audit_log.jsonl"
    log_path.write_text(
        json.dumps({"event": "ignored", "incomplete": True})
        + "\n"
        + json.dumps({"authorization_record": AUTHORIZATION_RECORD.__dict__})
        + "\n",
        encoding="utf-8",
    )

    assert load_authorization_record(log_path) == AUTHORIZATION_RECORD


def test_authorization_record_renders_to_docx_and_pdf(tmp_path: Path) -> None:
    docx_path = tmp_path / "Day32_Authorization_Record_Test.docx"
    pdf_path = tmp_path / "Day32_Authorization_Record_Test.pdf"

    document = Document()
    add_authorization_record_docx(document, AUTHORIZATION_RECORD)
    document.save(docx_path)
    docx_text = "\n".join(
        paragraph.text for paragraph in Document(docx_path).paragraphs
    )
    table_text = "\n".join(
        cell.text for table in Document(docx_path).tables for row in table.rows for cell in row.cells
    )
    assert "Appendix: Authorization & Verification Record" in docx_text
    assert AUTHORIZATION_RECORD.target_domain in table_text
    assert "VERIFIED" in table_text

    pdf = SentinelReportPDF()
    pdf.set_compression(False)
    pdf.add_authorization_record_pdf(AUTHORIZATION_RECORD)
    pdf.output(pdf_path)
    pdf_bytes = pdf_path.read_bytes()
    assert b"Authorization & Verification Record" in pdf_bytes
    assert b"Suraj Yadav" in pdf_bytes
    assert b"VERIFIED" in pdf_bytes