"""DOCX rendering helpers for active dynamic-scan findings."""

from dataclasses import asdict, is_dataclass
from typing import Any, Iterable

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor


def _finding_value(finding: Any, name: str, default: str = "") -> str:
    if is_dataclass(finding):
        return str(getattr(finding, name, default) or default)
    if isinstance(finding, dict):
        return str(finding.get(name, default) or default)
    return str(getattr(finding, name, default) or default)


def set_cell_background(cell: Any, fill_hex: str) -> None:
    """Set a table cell's background color."""
    tc_pr = cell._element.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill_hex)
    tc_pr.append(shading)


def add_finding_card_docx(
    doc: Document, finding: Any, finding_type: str = "PASSIVE"
) -> None:
    """Render a finding with distinct active and passive visual treatment."""
    table = doc.add_table(rows=2, cols=1)
    table.style = "Table Grid"
    header = table.cell(0, 0)
    active = finding_type.upper() == "ACTIVE"
    set_cell_background(header, "D9534F" if active else "FCE4B2")
    paragraph = header.paragraphs[0]
    tag_run = paragraph.add_run(
        f" {'CONFIRMED EXPLOITABLE' if active else 'POTENTIAL FINDING'} "
    )
    tag_run.bold = True
    title_run = paragraph.add_run(
        f"| {_finding_value(finding, 'title', 'Untitled finding')}"
    )
    title_run.bold = True
    if active:
        tag_run.font.color.rgb = RGBColor(255, 255, 255)
        title_run.font.color.rgb = RGBColor(255, 255, 255)

    body = table.cell(1, 0).paragraphs[0]
    body.add_run(
        f"Severity: {_finding_value(finding, 'severity', 'Medium')}\n"
    ).bold = True
    body.add_run(
        f"Target: {_finding_value(finding, 'target', _finding_value(finding, 'target_url', 'N/A'))}\n"
    )
    body.add_run(f"Description: {_finding_value(finding, 'description')}\n")
    body.add_run(f"Confidence: {_finding_value(finding, 'confidence', 'N/A')}\n")
    body.add_run(f"Associated CVE: {_finding_value(finding, 'cve_id', 'N/A')}\n")
    body.add_run(f"Remediation: {_finding_value(finding, 'remediation', 'N/A')}")
    doc.add_paragraph()


def add_active_findings_section(
    doc: Document, active_findings: Iterable[Any]
) -> None:
    """Render the dedicated confirmed active-testing section in a DOCX."""
    findings = list(active_findings)
    doc.add_heading("Active Testing Findings", level=1)

    notice = doc.add_paragraph()
    notice.add_run(
        "Notice: The following vulnerabilities were identified via active "
        "dynamic testing. These represent confirmed/validated findings "
        "against the authorized target."
    ).italic = True

    if not findings:
        doc.add_paragraph("No active testing vulnerabilities were detected.")
        return

    for finding in findings:
        add_finding_card_docx(doc, finding, finding_type="ACTIVE")

        for evidence_name, label in (
            ("evidence_request", "Evidence (HTTP Request):"),
            ("evidence_response", "Evidence (HTTP Response):"),
        ):
            evidence = _finding_value(finding, evidence_name)
            if evidence:
                doc.add_heading(label, level=3)
                code = doc.add_paragraph()
                code.add_run(evidence).font.name = "Courier New"

        doc.add_paragraph()


def add_authorization_record_docx(doc: Document, auth_record: Any) -> None:
    """Append the authorization and ownership verification record."""
    doc.add_page_break()
    doc.add_heading("Appendix: Authorization & Verification Record", level=1)

    disclaimer = doc.add_paragraph()
    disclaimer.add_run(
        "Audit Trail Disclaimer: This section provides an immutable record of "
        "the authorization and ownership verification performed prior to "
        "conducting active testing."
    ).italic = True

    table = doc.add_table(rows=6, cols=2)
    table.style = "Table Grid"
    records = [
        ("Target Domain / IP:", _finding_value(auth_record, "target_domain", "N/A")),
        ("Authorized By (Attestation):", _finding_value(auth_record, "attestation_user", "N/A")),
        ("Timestamp (UTC):", _finding_value(auth_record, "attestation_timestamp", "N/A")),
        ("Verification Method:", _finding_value(auth_record, "verification_method", "N/A")),
        ("Verification Status:", _finding_value(auth_record, "verification_status", "UNVERIFIED")),
        ("Token Used:", _finding_value(auth_record, "token_used", "N/A")),
    ]
    for row, (label, value) in zip(table.rows, records):
        row.cells[0].paragraphs[0].add_run(label).bold = True
        value_run = row.cells[1].paragraphs[0].add_run(value)
        if label == "Verification Status:":
            value_run.font.color.rgb = RGBColor(
                0, 150, 0 if value.upper() == "VERIFIED" else 0
            )
            if value.upper() != "VERIFIED":
                value_run.font.color.rgb = RGBColor(200, 0, 0)
    doc.add_paragraph()