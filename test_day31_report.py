"""Smoke test for the Day 31 active findings report sections."""

from pathlib import Path

from docx import Document

from reports.docx_generator import add_active_findings_section
from reports.pdf_generator import SentinelReportPDF
from reports.schema import ActiveFinding


ACTIVE_FINDING = ActiveFinding(
    id="ACT-001",
    title="Cross-Site Scripting (Reflected)",
    severity="High",
    confidence="High",
    target_url="http://localhost:3000/search?q=test",
    description="Parameter 'q' is reflected without proper sanitization.",
    cve_id="CVE-2023-XXXX",
    evidence_request="GET /search?q=<script>alert(1)</script> HTTP/1.1",
)


def test_active_findings_render_to_docx_and_pdf(tmp_path: Path) -> None:
    docx_path = tmp_path / "Day31_Active_Report_Test.docx"
    pdf_path = tmp_path / "Day31_Active_Report_Test.pdf"

    document = Document()
    document.add_heading("Passive Findings", level=1)
    document.add_paragraph("Existing passive findings remain in the report.")
    add_active_findings_section(document, [ACTIVE_FINDING])
    document.save(docx_path)

    docx_text = "\n".join(
        paragraph.text for paragraph in Document(docx_path).paragraphs
    )
    assert "Passive Findings" in docx_text
    assert "Active Testing Findings" in docx_text
    assert "Cross-Site Scripting (Reflected)" in docx_text

    pdf = SentinelReportPDF()
    pdf.set_compression(False)
    pdf.add_page()
    pdf.set_font("Helvetica", size=10)
    pdf.cell(0, 8, "Passive Findings")
    pdf.add_active_findings_section([ACTIVE_FINDING])
    pdf.output(pdf_path)

    pdf_bytes = pdf_path.read_bytes()
    assert b"Active Testing Findings" in pdf_bytes
    assert b"Cross-Site Scripting" in pdf_bytes